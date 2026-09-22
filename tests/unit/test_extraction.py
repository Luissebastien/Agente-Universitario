import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock

from database import extraction_repository as extraction_repo
from database import ingestion_repository as ingestion_repo
from database.db import connect
from extraction.extract import Extraction, ExtractionError
from extraction.ocr import OcrEngine, OcrResult
from ingestion.models import ResourceDescriptor, ResourceVersion
from ingestion.storage import FilesystemStorage, StorageError, StorageNotFoundError


class _FakeOcrEngine(OcrEngine):
    name = "fake"
    version = "1"

    def ocr(self, image_bytes: bytes) -> OcrResult:
        return OcrResult(text="ocr text", metadata={})


def _descriptor(name: str, mimetype: str | None) -> ResourceDescriptor:
    return ResourceDescriptor(
        origin="moodle",
        source_type="file",
        external_reference=f"https://campusvirtual.intec.edu.do/webservice/pluginfile.php/1/{name}",
        course_id=None,
        section_id=None,
        module_id=None,
        name=name,
        source_url=f"https://campusvirtual.intec.edu.do/webservice/pluginfile.php/1/{name}",
        mimetype=mimetype,
    )


class ExtractionTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.conn = connect(":memory:")
        self._tmp = tempfile.TemporaryDirectory()
        self.storage = FilesystemStorage(Path(self._tmp.name) / "originals")
        self.extraction = Extraction(self.storage, self.conn)

    def tearDown(self) -> None:
        self.conn.close()
        self._tmp.cleanup()

    def _ingest(self, name: str, mimetype: str | None, content: bytes) -> ResourceVersion:
        resource = ingestion_repo.upsert_resource(self.conn, _descriptor(name, mimetype))
        storage_ref = self.storage.store(content, "irrelevant-in-these-tests")
        return ingestion_repo.insert_resource_version(
            self.conn,
            ResourceVersion(
                id=None,
                resource_id=resource.id,
                version_number=1,
                content_hash="irrelevant-in-these-tests",
                storage_ref=storage_ref,
                size_bytes=len(content),
                mimetype=mimetype,
                ingested_at="2026-01-01T00:00:00+00:00",
            ),
        )


class SupportedFormatTests(ExtractionTestCase):
    def test_extracts_plain_text(self) -> None:
        version = self._ingest("notes.txt", "text/plain", b"hello world")

        doc = self.extraction.extract(version.id)

        self.assertEqual(doc.status, "done")
        self.assertEqual(doc.extracted_text, "hello world")
        self.assertEqual(doc.extractor_name, "plain_text")
        self.assertIsNone(doc.error_reason)

    def test_extracts_html_visible_text(self) -> None:
        version = self._ingest("page.html", "text/html", b"<p>Hola</p>")

        doc = self.extraction.extract(version.id)

        self.assertEqual(doc.status, "done")
        self.assertEqual(doc.extracted_text, "Hola")
        self.assertEqual(doc.extractor_name, "html_text")

    def test_falls_back_to_extension_when_mimetype_is_none(self) -> None:
        # Mirrors a real observed case: Moodle sometimes reports no mimetype at all.
        version = self._ingest("page.html", None, b"<p>Sin mimetype</p>")

        doc = self.extraction.extract(version.id)

        self.assertEqual(doc.status, "done")
        self.assertEqual(doc.extractor_name, "html_text")

    def test_persists_the_attempt(self) -> None:
        version = self._ingest("notes.txt", "text/plain", b"hello world")
        doc = self.extraction.extract(version.id)

        stored = extraction_repo.get_extracted_documents(self.conn, version.id)
        self.assertEqual(len(stored), 1)
        self.assertEqual(stored[0].id, doc.id)


class UnsupportedFormatTests(ExtractionTestCase):
    """.doc (legacy binary Word) is deliberately out of MVP scope - the OCR
    benchmark found no viable local extraction strategy for it - so it's a
    realistic, still-unsupported fixture (unlike .pptx/.docx/etc, which this
    phase now supports)."""

    def test_unsupported_mimetype_fails_explicitly_without_raising(self) -> None:
        version = self._ingest("notes.doc", "application/msword", b"\x00\x01binary")

        doc = self.extraction.extract(version.id)

        self.assertEqual(doc.status, "failed")
        self.assertIsNone(doc.extracted_text)
        self.assertIn("unsupported format", doc.error_reason)

    def test_unsupported_format_does_not_delete_the_original(self) -> None:
        version = self._ingest("notes.doc", "application/msword", b"\x00\x01binary")
        self.extraction.extract(version.id)

        self.assertTrue(self.storage.exists(version.storage_ref))
        self.assertEqual(self.storage.retrieve(version.storage_ref), b"\x00\x01binary")


class ContentFailureTests(ExtractionTestCase):
    def test_undecodable_declared_text_fails_explicitly(self) -> None:
        version = self._ingest("mislabeled.txt", "text/plain", b"\xff\xfe\x00\x01")

        doc = self.extraction.extract(version.id)

        self.assertEqual(doc.status, "failed")
        self.assertIn("UnicodeDecodeError", doc.error_reason)

    def test_storage_retrieval_failure_fails_explicitly_without_raising(self) -> None:
        resource = ingestion_repo.upsert_resource(self.conn, _descriptor("gone.txt", "text/plain"))
        version = ingestion_repo.insert_resource_version(
            self.conn,
            ResourceVersion(
                id=None,
                resource_id=resource.id,
                version_number=1,
                content_hash="never-stored",
                storage_ref="never-stored",
                size_bytes=0,
                mimetype="text/plain",
                ingested_at="2026-01-01T00:00:00+00:00",
            ),
        )

        doc = self.extraction.extract(version.id)

        self.assertEqual(doc.status, "failed")
        self.assertIn("retrieve", doc.error_reason)


class InvalidRequestTests(ExtractionTestCase):
    def test_unknown_resource_version_raises(self) -> None:
        with self.assertRaises(ExtractionError):
            self.extraction.extract(999999)

    def test_invalid_depth_raises(self) -> None:
        version = self._ingest("notes.txt", "text/plain", b"hi")
        with self.assertRaises(ExtractionError):
            self.extraction.extract(version.id, depth="ultra")


class AppendOnlyTests(ExtractionTestCase):
    def test_re_extracting_creates_a_new_attempt_not_a_dedup(self) -> None:
        version = self._ingest("notes.txt", "text/plain", b"hello world")

        first = self.extraction.extract(version.id)
        second = self.extraction.extract(version.id)

        self.assertNotEqual(first.id, second.id)
        attempts = extraction_repo.get_extracted_documents(self.conn, version.id)
        self.assertEqual(len(attempts), 2)


class DepthTests(ExtractionTestCase):
    def test_depth_is_recorded_on_the_document(self) -> None:
        version = self._ingest("notes.txt", "text/plain", b"hello world")

        basic = self.extraction.extract(version.id, depth="basic")
        deep = self.extraction.extract(version.id, depth="deep")

        self.assertEqual(basic.depth, "basic")
        self.assertEqual(deep.depth, "deep")


class NewFormatsThroughExtractionTests(ExtractionTestCase):
    """End-to-end through Extraction.extract() (not just the extractor
    classes in isolation) for every MVP format, confirming the wiring in
    build_default_extractors() actually works - not just each extractor on
    its own."""

    def setUp(self) -> None:
        super().setUp()
        self.extraction_with_ocr = Extraction(self.storage, self.conn, ocr_engine=_FakeOcrEngine())

    def _docx_bytes(self) -> bytes:
        import docx

        doc = docx.Document()
        doc.add_paragraph("Contenido real del documento.")
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def _xlsx_bytes(self) -> bytes:
        import openpyxl

        wb = openpyxl.Workbook()
        wb.active.append(["Curso", "Nota"])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _png_bytes(self) -> bytes:
        from PIL import Image

        im = Image.new("RGB", (40, 20), "white")
        buf = io.BytesIO()
        im.save(buf, format="PNG")
        return buf.getvalue()

    def test_docx_through_full_pipeline(self) -> None:
        version = self._ingest(
            "notas.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            self._docx_bytes(),
        )
        doc = self.extraction_with_ocr.extract(version.id)
        self.assertEqual(doc.status, "done")
        self.assertIn("Contenido real", doc.extracted_text)
        self.assertEqual(doc.extractor_name, "docx")

    def test_xlsx_through_full_pipeline(self) -> None:
        version = self._ingest(
            "notas.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            self._xlsx_bytes(),
        )
        doc = self.extraction_with_ocr.extract(version.id)
        self.assertEqual(doc.status, "done")
        self.assertIn("Curso | Nota", doc.extracted_text)
        self.assertEqual(doc.extractor_name, "xlsx")

    def test_image_through_full_pipeline_uses_ocr(self) -> None:
        version = self._ingest("captura.png", "image/png", self._png_bytes())
        doc = self.extraction_with_ocr.extract(version.id)
        self.assertEqual(doc.status, "done")
        self.assertEqual(doc.extracted_text, "ocr text")
        self.assertEqual(doc.extractor_name, "image_ocr")

    def test_image_without_ocr_engine_is_unsupported_not_a_crash(self) -> None:
        extraction_no_ocr = Extraction(self.storage, self.conn, ocr_engine=None)
        version = self._ingest("captura.png", "image/png", self._png_bytes())
        doc = extraction_no_ocr.extract(version.id)
        self.assertEqual(doc.status, "failed")
        self.assertIn("unsupported format", doc.error_reason)

    def test_doc_legacy_remains_unsupported(self) -> None:
        # DOC is explicitly out of MVP scope - must still degrade safely.
        version = self._ingest("viejo.doc", "application/msword", b"\x00\x01legacy")
        doc = self.extraction_with_ocr.extract(version.id)
        self.assertEqual(doc.status, "failed")
        self.assertIn("unsupported format", doc.error_reason)
        self.assertTrue(self.storage.exists(version.storage_ref))  # original preserved


class DepthDoesNotChangeExtractorSelectionTests(ExtractionTestCase):
    def test_basic_and_deep_use_the_same_extractor_for_docx(self) -> None:
        extraction = Extraction(self.storage, self.conn, ocr_engine=_FakeOcrEngine())
        import docx

        d = docx.Document()
        d.add_paragraph("Igual en ambas profundidades.")
        buf = io.BytesIO()
        d.save(buf)
        version = self._ingest(
            "x.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", buf.getvalue()
        )

        basic = extraction.extract(version.id, depth="basic")
        deep = extraction.extract(version.id, depth="deep")

        self.assertEqual(basic.extractor_name, deep.extractor_name)
        self.assertEqual(basic.extracted_text, deep.extracted_text)


if __name__ == "__main__":
    unittest.main()
