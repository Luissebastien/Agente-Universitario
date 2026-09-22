from __future__ import annotations

import abc
import io
from dataclasses import dataclass, field
from html.parser import HTMLParser

from extraction.ocr import OcrEngine, deskew, extract_ooxml_slide_text, ocr_embedded_images

_PLAIN_TEXT_MIMETYPES = {"text/plain", "text/uri-list"}
_PLAIN_TEXT_EXTENSIONS = {".txt"}

_HTML_MIMETYPES = {"text/html"}
_HTML_EXTENSIONS = {".html", ".htm"}

# Elements whose contents are not visible document text.
_SKIP_TAGS = {"script", "style"}


@dataclass
class ExtractionResult:
    """What a successful Extractor.extract() call produces."""

    text: str
    metadata: dict = field(default_factory=dict)


class Extractor(abc.ABC):
    """A deterministic converter from raw bytes of a known format to text.

    No LLM, no interpretation of meaning - purely mechanical decoding/parsing.
    Selection is by declared mimetype with a filename-extension fallback,
    mirroring the same "extension + Moodle-declared MIME" classification
    already used elsewhere in this project rather than sniffing content.
    """

    name: str
    version: str

    @abc.abstractmethod
    def supports(self, mimetype: str | None, extension: str) -> bool: ...

    @abc.abstractmethod
    def extract(self, data: bytes) -> ExtractionResult: ...


class PlainTextExtractor(Extractor):
    """Decodes UTF-8 plain text content, including preserved URL references."""

    name = "plain_text"
    version = "1"

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return mimetype in _PLAIN_TEXT_MIMETYPES or extension in _PLAIN_TEXT_EXTENSIONS

    def extract(self, data: bytes) -> ExtractionResult:
        # Strict decode: undecodable bytes mean the content isn't actually
        # what it claimed to be, which should surface as an explicit failure
        # (DEC-048) rather than be silently mangled with errors="replace".
        text = data.decode("utf-8")
        return ExtractionResult(text=text, metadata={"char_count": len(text)})


class _VisibleTextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._chunks: list[str] = []
        self._skip_depth = 0
        self.tag_count = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        self.tag_count += 1
        if tag in _SKIP_TAGS:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIP_TAGS and self._skip_depth > 0:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0:
            stripped = data.strip()
            if stripped:
                self._chunks.append(stripped)

    def get_text(self) -> str:
        return "\n".join(self._chunks)


class HtmlExtractor(Extractor):
    """Strips markup to plain visible text using the stdlib HTML parser.

    Deliberately does not resolve/fetch anything referenced by the HTML
    (images, links, scripts) - that would be crawling, out of scope here.
    """

    name = "html_text"
    version = "1"

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return mimetype in _HTML_MIMETYPES or extension in _HTML_EXTENSIONS

    def extract(self, data: bytes) -> ExtractionResult:
        raw = data.decode("utf-8")
        parser = _VisibleTextHTMLParser()
        parser.feed(raw)
        text = parser.get_text()
        return ExtractionResult(
            text=text,
            metadata={
                "char_count": len(text),
                "raw_char_count": len(raw),
                "tag_count": parser.tag_count,
            },
        )


MIN_NATIVE_CHARS_PER_PAGE = 50


def is_native_text_sufficient(pages_text: list[str], min_chars_per_page: int = MIN_NATIVE_CHARS_PER_PAGE) -> bool:
    """Deterministic, standalone (no PDF library involved) criterion for
    whether a PDF's native text layer is usable.

    Uses the MEDIAN page's character count, not the mean or the total: a
    handful of image-only pages in an otherwise text-heavy PDF shouldn't
    wrongly trigger OCR of the whole document, and a handful of text-heavy
    pages (e.g. a cover/title page) shouldn't mask an otherwise-empty scan.
    The 50-chars-per-page threshold and the median choice were validated
    empirically in the OCR benchmark against 253 real PDFs (INTEC + public):
    it correctly separated documents with a usable text layer (~97%) from
    the handful that were genuinely image-only, with no observed false
    negatives in that corpus.
    """
    if not pages_text:
        return False
    counts = sorted(len(p) for p in pages_text)
    median = counts[len(counts) // 2]
    return median >= min_chars_per_page


class PdfExtractor(Extractor):
    """Hybrid PDF extraction: native text (pypdf) first, OCR only as a
    fallback when the native text layer is insufficient (see
    is_native_text_sufficient). This is the main cost-control mechanism -
    OCR is comparatively expensive (seconds per page vs milliseconds for
    native extraction) and most real PDFs already have usable text.

    ocr_engine may be None (no OCR configured). In that case, an
    insufficient-text PDF still completes successfully rather than failing -
    per DEC-048's own example ("scanned PDF without OCR -> extraction may
    complete with an explicit `scanned` flag and no extracted text") - with
    metadata explicitly flagging that OCR was not available, so a future
    reprocessing pass can find and retry it.
    """

    name = "pdf_hybrid"
    version = "1"

    def __init__(self, ocr_engine: OcrEngine | None = None) -> None:
        self._ocr_engine = ocr_engine

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return mimetype == "application/pdf" or extension == ".pdf"

    def extract(self, data: bytes) -> ExtractionResult:
        import pypdf

        reader = pypdf.PdfReader(io.BytesIO(data))
        pages_text = [(page.extract_text() or "") for page in reader.pages]
        n_pages = len(pages_text)
        native_chars = sum(len(p) for p in pages_text)
        sufficient = is_native_text_sufficient(pages_text)

        if sufficient:
            return ExtractionResult(
                text="\n".join(pages_text),
                metadata={"method": "native", "pages": n_pages, "native_chars": native_chars,
                          "sufficient_text": True},
            )

        if self._ocr_engine is None:
            return ExtractionResult(
                text="",
                metadata={"method": "native", "pages": n_pages, "native_chars": native_chars,
                          "sufficient_text": False, "scanned": True, "ocr_available": False},
            )

        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(data)
        try:
            ocr_texts = []
            for page in pdf:
                bitmap = page.render(scale=200 / 72)  # ~200 DPI, matches the OCR benchmark
                pil_image = bitmap.to_pil()
                buf = io.BytesIO()
                pil_image.save(buf, format="PNG")
                deskewed = deskew(buf.getvalue())
                ocr_texts.append(self._ocr_engine.ocr(deskewed).text)
        finally:
            pdf.close()

        return ExtractionResult(
            text="\n".join(ocr_texts),
            metadata={
                "method": "ocr", "pages": n_pages, "native_chars": native_chars,
                "sufficient_text": False, "scanned": True, "ocr_available": True,
                "ocr_engine": self._ocr_engine.name, "ocr_engine_version": self._ocr_engine.version,
                "preprocessing": "deskew",
            },
        )


_IMAGE_MIMETYPES = {"image/png", "image/jpeg", "image/jpg", "image/gif", "image/webp", "image/bmp", "image/tiff"}
_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff", ".tif"}


class ImageExtractor(Extractor):
    """Standalone image -> text via OCR (deskew + the configured OcrEngine).

    Unlike the other extractors, this one is only ever registered when an
    OcrEngine is actually available (see build_default_extractors) - an
    image with no OCR engine configured isn't a "done, empty" result the way
    a scanned PDF is, it's simply not a format we can currently process, so
    it correctly falls through to the existing unsupported-format path.
    """

    name = "image_ocr"
    version = "1"

    def __init__(self, ocr_engine: OcrEngine) -> None:
        self._ocr_engine = ocr_engine

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return mimetype in _IMAGE_MIMETYPES or extension in _IMAGE_EXTENSIONS

    def extract(self, data: bytes) -> ExtractionResult:
        deskewed = deskew(data)
        result = self._ocr_engine.ocr(deskewed)
        return ExtractionResult(
            text=result.text,
            metadata={
                "method": "ocr", "ocr_engine": self._ocr_engine.name,
                "ocr_engine_version": self._ocr_engine.version, "preprocessing": "deskew",
                **result.metadata,
            },
        )


class DocxExtractor(Extractor):
    """DOCX text/table extraction (python-docx) plus, when an OcrEngine is
    configured, OCR of embedded images appended to the text - some real DOCX
    tutorials are almost entirely screenshots with minimal native text (see
    the OCR benchmark: one real INTEC tutorial had 1,470 chars of text but 17
    embedded images), so skipping them would lose most of the document.
    """

    name = "docx"
    version = "1"

    def __init__(self, ocr_engine: OcrEngine | None = None) -> None:
        self._ocr_engine = ocr_engine

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return (
            extension == ".docx"
            or mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    def extract(self, data: bytes) -> ExtractionResult:
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)

        metadata = {
            "method": "native", "paragraphs": len(document.paragraphs),
            "tables": len(document.tables), "images_embedded": len(document.inline_shapes),
        }
        if self._ocr_engine is not None:
            embedded_texts, embedded_meta = ocr_embedded_images(data, "word/media/", self._ocr_engine)
            if embedded_texts:
                text = text + "\n\n" + "\n\n".join(embedded_texts)
            metadata["embedded_images_ocr"] = embedded_meta
        else:
            metadata["embedded_images_ocr"] = None

        return ExtractionResult(text=text, metadata=metadata)


class PptxExtractor(Extractor):
    """PPTX/PPSX text extraction via a manual OOXML reader (see
    extraction.ocr.extract_ooxml_slide_text) rather than python-pptx.

    python-pptx handles real .pptx files correctly but rejects every real
    .ppsx file tested in the OCR benchmark (it validates the internal
    content-type and only accepts 'presentationml.presentation'). The manual
    zip+XML reader doesn't care about the declared content-type - it reads
    the same slide XML structure either format actually uses - so one code
    path covers both formats reliably, confirmed against real files.

    Also OCRs embedded images when an OcrEngine is configured, same
    reasoning as DocxExtractor.
    """

    name = "pptx"
    version = "1"

    def __init__(self, ocr_engine: OcrEngine | None = None) -> None:
        self._ocr_engine = ocr_engine

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return extension in (".pptx", ".ppsx") or mimetype in (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "application/vnd.openxmlformats-officedocument.presentationml.slideshow",
        )

    def extract(self, data: bytes) -> ExtractionResult:
        text, n_slides = extract_ooxml_slide_text(data)
        metadata = {"method": "manual_ooxml", "slides": n_slides}

        if self._ocr_engine is not None:
            embedded_texts, embedded_meta = ocr_embedded_images(data, "ppt/media/", self._ocr_engine)
            if embedded_texts:
                text = text + "\n\n" + "\n\n".join(embedded_texts)
            metadata["embedded_images_ocr"] = embedded_meta
        else:
            metadata["embedded_images_ocr"] = None

        return ExtractionResult(text=text, metadata=metadata)


class XlsxExtractor(Extractor):
    """XLSX structured extraction (openpyxl): one line per non-empty row,
    cells joined with ' | ', prefixed by a '# <sheet name>' header per
    sheet - readable and keeps row/column structure legible instead of a
    flat character dump."""

    name = "xlsx"
    version = "1"

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return extension in (".xlsx", ".xlsm") or mimetype in (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        )

    def extract(self, data: bytes) -> ExtractionResult:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        lines = []
        n_cells = 0
        for ws in wb.worksheets:
            lines.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    n_cells += len(cells)
                    lines.append(" | ".join(cells))
        return ExtractionResult(
            text="\n".join(lines),
            metadata={"method": "native", "sheets": len(wb.worksheets), "cells": n_cells},
        )


class XlsExtractor(Extractor):
    """Legacy binary .xls extraction (xlrd - the only viable extractor found
    for this format in the OCR benchmark, ~88% success on real files; xlrd
    2.x deliberately dropped .xlsx support, so it is used for .xls only)."""

    name = "xls"
    version = "1"

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return extension == ".xls" or mimetype == "application/vnd.ms-excel"

    def extract(self, data: bytes) -> ExtractionResult:
        import xlrd

        wb = xlrd.open_workbook(file_contents=data)
        lines = []
        n_cells = 0
        for sheet in wb.sheets():
            lines.append(f"# {sheet.name}")
            for r in range(sheet.nrows):
                cells = [
                    str(sheet.cell_value(r, c))
                    for c in range(sheet.ncols)
                    if sheet.cell_value(r, c) not in ("", None)
                ]
                if cells:
                    n_cells += len(cells)
                    lines.append(" | ".join(cells))
        return ExtractionResult(
            text="\n".join(lines),
            metadata={"method": "native", "sheets": wb.nsheets, "cells": n_cells},
        )


class OdtExtractor(Extractor):
    """ODT (OpenDocument Text) extraction via odfpy."""

    name = "odt"
    version = "1"

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return extension == ".odt" or mimetype == "application/vnd.oasis.opendocument.text"

    def extract(self, data: bytes) -> ExtractionResult:
        from odf import teletype as odf_teletype
        from odf import text as odf_text
        from odf.opendocument import load as odf_load

        doc = odf_load(io.BytesIO(data))
        paragraphs = doc.getElementsByType(odf_text.P)
        parts = [odf_teletype.extractText(p) for p in paragraphs]
        return ExtractionResult(
            text="\n".join(parts), metadata={"method": "native", "paragraphs": len(parts)}
        )


class OdsExtractor(Extractor):
    """ODS (OpenDocument Spreadsheet) structured extraction via odfpy - same
    readable row/cell layout as XlsxExtractor."""

    name = "ods"
    version = "1"

    def supports(self, mimetype: str | None, extension: str) -> bool:
        return extension == ".ods" or mimetype == "application/vnd.oasis.opendocument.spreadsheet"

    def extract(self, data: bytes) -> ExtractionResult:
        from odf import table as odf_table
        from odf import teletype as odf_teletype
        from odf import text as odf_text
        from odf.opendocument import load as odf_load

        doc = odf_load(io.BytesIO(data))
        tables = doc.getElementsByType(odf_table.Table)
        lines = []
        n_cells = 0
        for tbl in tables:
            lines.append(f"# {tbl.getAttribute('name')}")
            for row in tbl.getElementsByType(odf_table.TableRow):
                cells_text = []
                for cell in row.getElementsByType(odf_table.TableCell):
                    value = " ".join(
                        odf_teletype.extractText(p) for p in cell.getElementsByType(odf_text.P)
                    )
                    if value:
                        cells_text.append(value)
                if cells_text:
                    n_cells += len(cells_text)
                    lines.append(" | ".join(cells_text))
        return ExtractionResult(
            text="\n".join(lines), metadata={"method": "native", "tables": len(tables), "cells": n_cells}
        )


def build_default_extractors(ocr_engine: OcrEngine | None = None) -> list[Extractor]:
    """The extractor list Extraction uses when none is explicitly injected.

    ImageExtractor is only included when ocr_engine is not None (see its
    docstring). Every other extractor is always included - PdfExtractor
    degrades gracefully without an OCR engine (DEC-048 "scanned" case), and
    the document extractors just skip embedded-image OCR when unavailable.
    """
    extractors: list[Extractor] = [
        PlainTextExtractor(),
        HtmlExtractor(),
        PdfExtractor(ocr_engine),
        DocxExtractor(ocr_engine),
        PptxExtractor(ocr_engine),
        XlsxExtractor(),
        XlsExtractor(),
        OdtExtractor(),
        OdsExtractor(),
    ]
    if ocr_engine is not None:
        extractors.append(ImageExtractor(ocr_engine))
    return extractors


# Ordered list of known extractors - first match wins. Kept as the default
# used when no ocr_engine is configured (backwards compatible with existing
# callers); build_default_extractors(ocr_engine) is what Extraction actually
# calls now - see extraction/extract.py.
DEFAULT_EXTRACTORS: list[Extractor] = build_default_extractors(None)
